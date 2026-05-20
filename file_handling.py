# import os

#def create_file(): 
# def write_file(value):
#     with open('text.xlsx', 'w') as f:
#         f.write(value)
# write_file("\nHello Data Science its Mondayy")
# write_file("\nHello Data Science its tuesdayy")
# def read_file():
#     with open('text.xlsx', 'r') as f:
#         content = f.read()
#         print(content)

# read_file()
# def append_file(value):
#     with open('text.xlsx', 'a') as f:
#         f.write(value)
# append_file("\nHello Data Science its wednesdayy")
# # read_file()
# def delete_file():
#     if os.path.exists('text.xlsx'):
#         os.remove('text.xlsx')

import openpyxl
def write_excel():
    wb = openpyxl.Workbook() 
    ws = wb.active
    ws['A1'] = 'Name'
    ws['B1'] = 'Age'
    ws['C1'] = 'City'
    ws['A2'] = 'Vishwash'
    ws['B2'] = 21
    ws['C2'] = 'Daman'
    wb.save("text.xlsx")

write_excel()
